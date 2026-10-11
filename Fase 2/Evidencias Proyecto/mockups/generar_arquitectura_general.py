# -*- coding: utf-8 -*-
"""Genera un diagrama técnico con identidad propia para la arquitectura general."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

BASE = Path(__file__).resolve().parent
OUT = BASE / "Arquitectura_general_Evento_Cultural_FINAL.pptx"
ICONS = BASE.parent / "ppt capsula" / "diagramas_d2" / "iconos" / "png"

NAVY = RGBColor(0x06, 0x16, 0x26)
NAVY_2 = RGBColor(0x0C, 0x26, 0x39)
STEEL = RGBColor(0x19, 0x3B, 0x52)
CYAN = RGBColor(0x27, 0xD3, 0xFF)
MINT = RGBColor(0x55, 0xE6, 0xB4)
YELLOW = RGBColor(0xFF, 0xD1, 0x66)
ORANGE = RGBColor(0xFF, 0x7A, 0x45)
WHITE = RGBColor(0xF5, 0xFB, 0xFD)
ICE = RGBColor(0xE9, 0xF4, 0xF7)
INK = RGBColor(0x0A, 0x20, 0x31)
MUTED = RGBColor(0x5B, 0x70, 0x7C)
GRID = RGBColor(0x18, 0x36, 0x4A)
LINE = RGBColor(0xB9, 0xD1, 0xD9)
HEAD = "Arial"
MONO = "Courier New"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(bg):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = bg
    return s


def rect(s, x, y, w, h, fill, line_color=None, width=1.0, transparency=0):
    sh = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill; sh.fill.transparency = transparency
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color; sh.line.width = Pt(width)
    return sh


def circle(s, x, y, d, fill, line_color=None, width=1.0):
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line_color; sh.line.width = Pt(width)
    return sh


def text(s, x, y, w, h, value, size=12, color=INK, bold=False, font=HEAD,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    values = value if isinstance(value, list) else [value]
    for i, item in enumerate(values):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(2)
        r = p.add_run(); r.text = item
        r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb


def line(s, x1, y1, x2, y2, color, width=1.2):
    sh = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb = color; sh.line.width = Pt(width)
    return sh


def grid(s, dark=True):
    color = GRID if dark else RGBColor(0xD8, 0xE8, 0xED)
    for i in range(1, 27):
        x = i * 0.5
        line(s, x, 0, x, 7.5, color, 0.35)
    for i in range(1, 15):
        y = i * 0.5
        line(s, 0, y, 13.333, y, color, 0.35)


def icon(s, name, x, y, size=0.42, white=False):
    suffix = "_w" if white and name in {"angular", "githubactions", "nvidia", "supabase"} else ""
    path = ICONS / f"{name}{suffix}.png"
    if not path.exists():
        path = ICONS / f"{name}.png"
    if path.exists():
        s.shapes.add_picture(str(path), Inches(x), Inches(y), width=Inches(size), height=Inches(size))


def code_label(s, x, y, label, color=CYAN, w=1.32, light=False):
    fill = WHITE if light else NAVY
    rect(s, x, y, w, 0.27, fill, color, 0.8)
    text(s, x, y + 0.06, w, 0.14, label, 7.2, color, True, MONO, PP_ALIGN.CENTER)


def status(s, x, y, label, color, dark=False):
    circle(s, x, y + 0.04, 0.09, color)
    text(s, x + 0.16, y, 1.25, 0.18, label, 7.4, WHITE if dark else MUTED, True, MONO)


def node(s, x, y, w, h, tag, title, detail, accent, logo=None, sequence=None):
    rect(s, x, y, w, h, NAVY_2, accent, 1.35)
    if sequence:
        text(s, x + 0.2, y + 0.17, 0.46, 0.25, sequence, 11, accent, True, MONO)
    if logo:
        icon(s, logo, x + w - 0.58, y + 0.16, 0.36, white=True)
    text(s, x + 0.2, y + 0.55, w - 0.4, 0.25, tag, 7.5, accent, True, MONO)
    text(s, x + 0.2, y + 0.88, w - 0.4, 0.34, title, 14, WHITE, True)
    text(s, x + 0.2, y + 1.32, w - 0.4, h - 1.43, detail, 9.1, RGBColor(0xBA, 0xD0, 0xD9))


def footer(s, num, dark=True):
    color = RGBColor(0x74, 0x93, 0xA2) if dark else MUTED
    text(s, 0.55, 7.17, 4.4, 0.16, "ARQUITECTURA GENERAL / EVENTO CULTURAL", 7.2, color, True, MONO)
    text(s, 12.25, 7.17, 0.48, 0.16, f"{num:02d}", 7.2, color, True, MONO, PP_ALIGN.RIGHT)


# Diapositiva 1: plano técnico de la arquitectura.
s = add_slide(NAVY)
grid(s, True)
text(s, 0.55, 0.42, 2.9, 0.2, "SYSTEM BLUEPRINT / v2.0", 8.5, CYAN, True, MONO)
text(s, 0.55, 0.76, 8.4, 0.55, "Arquitectura general", 30, WHITE, True)
text(s, 0.58, 1.34, 7.5, 0.28, "Mapa técnico del flujo implementado y sus puntos de integración.", 11, RGBColor(0xA9, 0xC4, 0xCE))
code_label(s, 10.86, 0.53, "FASE 2", MINT, 0.86)
code_label(s, 11.84, 0.53, "ACTUAL", CYAN, 0.94)

text(s, 0.58, 1.92, 2.0, 0.2, "01 / ENTRADAS", 8.5, CYAN, True, MONO)
entries = [
    (2.30, "FUENTES PÚBLICAS", "22 agendas · HTTP / REST / Playwright", ORANGE, None, "22"),
    (3.45, "GITHUB ACTIONS", "cron 23:37 · pruebas · checkpoints", CYAN, "githubactions", None),
    (4.60, "NVIDIA NIM", "OCR selectivo · enriquecimiento · 40 RPM", MINT, "nvidia", None),
]
for y, title_, detail, accent, logo, count in entries:
    rect(s, 0.58, y, 3.15, 0.84, NAVY_2, accent, 1.1)
    if count:
        text(s, 0.76, y + 0.17, 0.48, 0.32, count, 17, accent, True, MONO)
    elif logo:
        icon(s, logo, 0.76, y + 0.2, 0.42, white=True)
    text(s, 1.38, y + 0.14, 2.08, 0.22, title_, 10.2, WHITE, True)
    text(s, 1.38, y + 0.45, 2.08, 0.19, detail, 7.8, RGBColor(0xA9, 0xC4, 0xCE))
    line(s, 3.73, y + 0.42, 4.22, y + 0.42, accent, 1.3)
    circle(s, 4.16, y + 0.36, 0.12, accent)

line(s, 4.22, 2.72, 4.22, 5.43, CYAN, 1.4)
text(s, 4.04, 5.60, 0.42, 0.18, "BUS", 7.5, CYAN, True, MONO, PP_ALIGN.CENTER)
text(s, 4.78, 1.92, 2.8, 0.2, "02 / PROCESAMIENTO Y ENTREGA", 8.5, CYAN, True, MONO)
node(s, 4.78, 2.30, 3.38, 1.65, "EXTRACT / NORMALIZE", "Scraper Python",
     "10 familias de conectores · calidad · checkpoints SQLite", ORANGE, "python", "A1")
node(s, 8.63, 2.30, 3.38, 1.65, "STORE / GOVERN", "Supabase",
     "PostgreSQL + RLS · catálogo canónico · procedencia privada", MINT, "supabase", "A2")
node(s, 4.78, 4.45, 3.38, 1.65, "PRESENT / EXPLORE", "Angular + Ionic",
     "Interfaz web y móvil · actualmente consume datos mock", YELLOW, "angular", "A3")
node(s, 8.63, 4.45, 3.38, 1.65, "DISCOVER / REDIRECT", "Usuarios",
     "Descubren, filtran y consultan la fuente oficial del evento", CYAN, None, "A4")

line(s, 4.22, 3.13, 4.78, 3.13, ORANGE, 1.8)
line(s, 8.16, 3.13, 8.63, 3.13, MINT, 1.8)
line(s, 10.32, 3.95, 10.32, 4.45, CYAN, 1.8)
line(s, 8.63, 5.27, 8.16, 5.27, YELLOW, 1.8)
circle(s, 4.16, 3.07, 0.12, ORANGE)
circle(s, 8.10, 3.07, 0.12, MINT)
circle(s, 10.26, 3.89, 0.12, CYAN)
circle(s, 8.10, 5.21, 0.12, YELLOW)

rect(s, 4.78, 6.42, 7.23, 0.43, NAVY_2, STEEL, 0.8)
status(s, 5.04, 6.54, "IMPLEMENTADO", MINT, True)
text(s, 6.48, 6.52, 5.22, 0.18, "PENDIENTE: sustituir mock del frontend por el catálogo real de Supabase.",
     8.2, RGBColor(0xB9, 0xD1, 0xD9), False, MONO)
footer(s, 1, True)


# Diapositiva 2: registro técnico, sin cuadrícula de tarjetas.
s = add_slide(WHITE)
grid(s, False)
text(s, 0.55, 0.42, 3.7, 0.2, "COMPONENT REGISTER / RESPONSABILIDADES", 8.5, MUTED, True, MONO)
text(s, 0.55, 0.78, 8.7, 0.52, "Inventario tecnológico", 29, INK, True)
text(s, 0.58, 1.34, 9.2, 0.26, "Lectura operativa del stack: qué resuelve cada componente y cuál es su estado.", 10.8, MUTED)
code_label(s, 11.28, 0.56, "8 COMPONENTES", CYAN, 1.50, True)

items = [
    ("01", "Python", "Scraping, normalización y pipeline", "python", "IMPLEMENTADO", MINT),
    ("02", "SQLite", "Checkpoints, caché y cola OCR", "sqlite", "IMPLEMENTADO", MINT),
    ("03", "NVIDIA NIM", "OCR y enriquecimiento selectivo · 40 RPM", "nvidia", "IMPLEMENTADO", MINT),
    ("04", "GitHub Actions", "Pruebas y ejecución diaria a las 23:37", "githubactions", "IMPLEMENTADO", MINT),
    ("05", "Supabase + PostgreSQL", "Catálogo, staging, RPC y RLS", "supabase", "IMPLEMENTADO", MINT),
    ("06", "Angular 22 + Ionic 9", "Interfaz web/móvil; hoy usa datos mock", "angular", "IMPLEMENTADO", MINT),
    ("07", "Vercel + Capacitor", "Despliegue web y empaquetado móvil", None, "PREVISTO", YELLOW),
    ("08", "Integración frontend–BD", "Sustituir datos mock por catálogo real", None, "PENDIENTE", ORANGE),
]
for i, (num, name, detail, logo, state, color) in enumerate(items):
    col, row = i // 4, i % 4
    x = 0.58 if col == 0 else 6.82
    y = 1.98 + row * 1.13
    line(s, x, y + 0.93, x + 5.82, y + 0.93, LINE, 0.7)
    text(s, x, y + 0.07, 0.5, 0.28, num, 13, color, True, MONO)
    circle(s, x + 0.62, y + 0.03, 0.52, NAVY_2)
    if logo:
        icon(s, logo, x + 0.73, y + 0.14, 0.30, white=True)
    else:
        text(s, x + 0.62, y + 0.20, 0.52, 0.14, "<>" if num == "08" else "DEP", 7, WHITE, True, MONO, PP_ALIGN.CENTER)
    text(s, x + 1.32, y + 0.02, 3.05, 0.27, name, 13, INK, True)
    text(s, x + 1.32, y + 0.39, 3.9, 0.23, detail, 9.2, MUTED)
    status(s, x + 4.45, y + 0.03, state, color, False)

rect(s, 0.58, 6.62, 12.16, 0.33, ICE, LINE, 0.6)
text(s, 0.79, 6.71, 11.75, 0.15,
     "PRINCIPIO DE DISEÑO  /  extracción, persistencia e interfaz permanecen desacopladas para facilitar mantenimiento y escalado.",
     7.9, INK, True, MONO)
footer(s, 2, False)

prs.save(OUT)
print(f"OK: {OUT}")
print(f"Diapositivas: {len(prs.slides)}")
